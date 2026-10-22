"""Word 解析（python-docx；.doc 先用 LibreOffice 转 .docx，docs/04 §3.1）。"""
from __future__ import annotations

import subprocess
import tempfile
from pathlib import Path
from typing import List

from loguru import logger

from app.core.constants import ErrorCode
from app.core.exceptions import BizError
from app.rag.parser.meta import extract_meta
from app.rag.types import PageContent, ParsedDocument, TextBlock

try:
    import docx  # python-docx
except ImportError:  # pragma: no cover
    docx = None  # type: ignore


class DocxParser:
    def parse(self, file_path: str) -> ParsedDocument:
        path = Path(file_path)
        if docx is None:
            raise BizError(ErrorCode.PARSE_FAILED, "Word 解析组件 python-docx 未安装")

        work_path = path
        if path.suffix.lower() == ".doc":
            work_path = self._convert_doc(path)

        try:
            document = docx.Document(str(work_path))
        except Exception as exc:
            raise BizError(ErrorCode.PARSE_FAILED, "文档解析失败，文件可能已损坏") from exc

        blocks: List[TextBlock] = []
        # 段落（保留标题层级）
        for para in document.paragraphs:
            text = (para.text or "").strip()
            if not text:
                continue
            level = None
            style_name = (para.style.name or "") if para.style else ""
            if style_name.startswith("Heading"):
                try:
                    level = int(style_name.split()[-1])
                except (ValueError, IndexError):
                    level = 1
            blocks.append(
                TextBlock(
                    type="title" if level else "text",
                    text=text,
                    page_no=1,
                    level=level,
                )
            )
        # 表格 → Markdown（带表头前缀，docs/04 §3.3）
        for tbl in document.tables:
            rows = [[c.text.strip() for c in row.cells] for row in tbl.rows]
            if not rows:
                continue
            header = rows[0]
            lines = ["【表】", "| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
            for row in rows[1:]:
                lines.append("| " + " | ".join(row) + " |")
            blocks.append(TextBlock(type="table", text="\n".join(lines), page_no=1))

        if not blocks:
            raise BizError(ErrorCode.PARSE_FAILED, "未能从文件中提取到有效文本")

        # Word 无固定页码概念：整篇视为 1 页（页码溯源对 PDF 为准）
        page = PageContent(page_no=1, blocks=blocks, char_count=sum(len(b.text) for b in blocks), used_ocr=False)
        full_text = "\n".join(b.text for b in blocks)
        meta = extract_meta(full_text)
        meta.setdefault("title", path.stem)
        return ParsedDocument(
            file_name=path.name,
            file_ext=path.suffix.lower(),
            pages=[page],
            parse_method="text",
            meta=meta,
            warnings=[],
        )

    def _convert_doc(self, path: Path) -> Path:
        """.doc → .docx（LibreOffice，容器内可选安装）。"""
        with tempfile.TemporaryDirectory() as tmp:
            try:
                subprocess.run(
                    ["soffice", "--headless", "--convert-to", "docx", "--outdir", tmp, str(path)],
                    check=True,
                    capture_output=True,
                    timeout=60,
                )
            except Exception as exc:
                logger.warning(".doc 转换失败: {}", exc)
                raise BizError(
                    ErrorCode.PARSE_FAILED,
                    "暂不支持解析 .doc 格式，请另存为 .docx 或 .pdf 后重新上传",
                ) from exc
            converted = Path(tmp) / (path.stem + ".docx")
            if not converted.exists():
                raise BizError(ErrorCode.PARSE_FAILED, "文档转换失败，请另存为 .docx 或 .pdf")
            target = path.with_suffix(".docx.tmp.docx")
            target.write_bytes(converted.read_bytes())
            return target


docx_parser = DocxParser()
