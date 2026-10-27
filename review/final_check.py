# -*- coding: utf-8 -*-
"""终检：quote_text 逐字一致性、page_no 完整性、状态分布、passed 备注留空规则"""
import openpyxl, re, collections

wb = openpyxl.load_workbook(r'成员1_知识点审核.xlsx', read_only=True)
ws = wb['知识点候选']
hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
I = {h: i for i, h in enumerate(hdr)}
rows = list(ws.iter_rows(min_row=2, values_only=True))
arts = {}
for r in wb['完整条文'].iter_rows(min_row=2, values_only=True):
    arts[(r[0], r[1])] = {'text': r[2] or '', 'page': r[4]}

def norm(s):
    return re.sub(r'[\s\u3000]', '', str(s or ''))

err = 0
for r in rows:
    a = arts.get((r[I['doc_title']], r[I['article_no']]))
    if not a:
        print('缺条文:', r[I['code']]); err += 1; continue
    if norm(r[I['quote_text']]) not in norm(a['text']):
        print('quote不符:', r[I['code']]); err += 1
    if r[I['page_no']] in (None, ''):
        print('缺页码:', r[I['code']], '(须留疑点备注)')
    if r[I['verify_status']] == 'passed' and r[I['review_note']] not in (None, ''):
        print('passed带备注:', r[I['code']]); err += 1

print('校验错误数:', err)
print('verify_status:', collections.Counter(r[I['verify_status']] for r in rows))
print('status:', collections.Counter(r[I['status']] for r in rows))
print('非passed行:')
for r in rows:
    if r[I['verify_status']] != 'passed':
        print(' ', r[I['code']], r[I['verify_status']], '|', r[I['review_note']])
print('doc_no空(应仅证券法355行):', collections.Counter(r[I['doc_title']] for r in rows if r[I['doc_no']] in (None, '')))
print('expiry_date空:', sum(1 for r in rows if r[I['expiry_date']] in (None, '')), '（无失效依据，留空正确）')
print('title为空:', sum(1 for r in rows if r[I['title']] in (None, '')))
print('总行数:', len(rows))
