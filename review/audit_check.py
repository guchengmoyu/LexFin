# -*- coding: utf-8 -*-
"""知识点候选表自动核对脚本：quote_text / page_no / 重复 / 结构质量"""
import openpyxl, re, json, collections

XLSX = r'成员1_知识点审核.xlsx'
wb = openpyxl.load_workbook(XLSX, read_only=True)

ws = wb['知识点候选']
hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
I = {h: i for i, h in enumerate(hdr)}
rows = list(ws.iter_rows(min_row=2, values_only=True))

# 完整条文
arts = {}
ws2 = wb['完整条文']
for r in ws2.iter_rows(min_row=2, values_only=True):
    doc, art, text, src, page, url = r[:6]
    arts[(doc, art)] = {'text': text or '', 'page': page, 'src': src}

def norm(s):
    if s is None:
        return ''
    s = str(s)
    s = s.replace('\u3000', '').replace(' ', '').replace('\t', '')
    s = s.replace('\r', '').replace('\n', '')
    return s

report = collections.defaultdict(list)

for r in rows:
    code = r[I['code']]
    doc = r[I['doc_title']]
    art = r[I['article_no']]
    q = r[I['quote_text']]
    content = r[I['content']] or ''
    title = r[I['title']] or ''
    page = r[I['page_no']]

    a = arts.get((doc, art))
    if a is None:
        report['missing_article'].append((code, doc, art))
        continue
    atext, apage = a['text'], a['page']

    # 1) quote_text 是否逐字出现在原文
    if norm(q) not in norm(atext):
        report['quote_mismatch'].append((code, doc, art, q, atext))

    # 2) page_no 与完整条文页码
    if page in (None, ''):
        report['page_missing'].append((code, doc, art, apage))
    elif str(page) != str(apage):
        report['page_mismatch'].append((code, str(page), str(apage), doc, art))

    # 3) content 中是否含原文引用核对（content 摘录的正文句子是否在原文中）
    body = content.split('适用条件与同条上下文：')[0]
    body = body.replace('根据《中华人民共和国证券法》', '').replace('根据《保险销售行为管理办法》', '')
    body = body.replace('《中华人民共和国证券法》', '本法').replace('《保险销售行为管理办法》', '本办法')
    body = re.sub(r'^第[一二三四五六七八九十百零]+条[：:]\s*', '', body.strip())
    nb = norm(body)
    na = norm(atext).replace('本法', '本法')
    # 允许 content 把“本法/本办法”替换为全称
    na_alt = norm(atext).replace('本法', '中华人民共和国证券法').replace('本办法', '保险销售行为管理办法')
    if nb and nb not in norm(atext) and nb not in na_alt:
        # 逐句检查
        sents = [s for s in re.split(r'[。；]', body) if len(norm(s)) > 6]
        bad = []
        for s in sents:
            ns = norm(s)
            if ns in norm(atext) or ns in na_alt:
                continue
            ns2 = ns.replace('中华人民共和国证券法', '本法').replace('保险销售行为管理办法', '本办法')
            if ns2 in norm(atext):
                continue
            bad.append(s)
        if bad:
            report['content_sentence_not_in_source'].append((code, doc, art, bad[:3]))

    # 4) 长度
    if len(norm(q)) > 260:
        report['quote_too_long'].append((code, len(norm(q)), doc, art))
    if len(norm(content)) > 700:
        report['content_too_long'].append((code, len(norm(content)), doc, art))

    # 5) 上下文依赖标记
    if re.search(r'前款|上款|下列情形|如下|本条第|上述规定', str(q)) and not re.search(r'下列|如下', str(q)):
        report['context_dependent'].append((code, doc, art, title))

    # 6) 标题质量：疑问句
    if not re.search(r'[？?]', title):
        report['title_not_question'].append((code, title))

    # 7) 立法目的/定义类
    if re.search(r'立法目的|为了.*制定|本(法|办法)所称|是指', title) or re.search(r'^为了.*制定(本法|《)', norm(q)):
        report['purpose_or_definition'].append((code, title, doc, art))

# 8) 重复检测：同 quote / 高相似 title
qmap = collections.defaultdict(list)
for r in rows:
    qmap[(r[I['doc_title']], norm(r[I['quote_text']]))].append(r[I['code']])
for k, v in qmap.items():
    if len(v) > 1:
        report['dup_same_quote'].append((k[0], v, k[1][:60]))

tmap = collections.defaultdict(list)
for r in rows:
    key = norm(re.sub(r'[？?。，,、\s]', '', str(r[I['title']])))
    tmap[key].append(r[I['code']])
for k, v in tmap.items():
    if len(v) > 1:
        report['dup_same_title'].append((v, k[:60]))

out = {k: v for k, v in report.items()}
with open(r'review/audit_report.json', 'w', encoding='utf-8') as f:
    json.dump(out, f, ensure_ascii=False, indent=1)

for k in sorted(out):
    print(f'== {k}: {len(out[k])}')
    for item in out[k][:12]:
        s = str(item)
        print('   ', s[:220])
