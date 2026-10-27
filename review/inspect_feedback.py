"""核对陈玉婷对《成员1_知识点审核》的三条反馈：逐条调出原文，验证其说法。"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "成员1_知识点审核_待修订.xlsx"

TARGETS = ["KP-INS-0381", "KP-SEC-0232",
           "KP-SEC-0064", "KP-SEC-0101", "KP-SEC-0269", "KP-SEC-0070"]

wb = load_workbook(XLSX, read_only=True, data_only=True)
ws = wb["知识点候选"]
rows = ws.iter_rows(values_only=True)
header = [str(c).strip() if c else "" for c in next(rows)]
col = {n: i for i, n in enumerate(header)}

print("列序:", " | ".join(f"{i}:{n}" for i, n in enumerate(header)))
print("=" * 78)

for excel_row, row in enumerate(rows, start=2):
    code = row[col["code"]]
    if code not in TARGETS:
        continue
    g = lambda k: (row[col[k]] if k in col and row[col[k]] is not None else "")
    print(f"\n### Excel 第 {excel_row} 行  |  {code}")
    print(f"  title        : {g('title')}")
    print(f"  doc_title    : {g('doc_title')}   article_no: {g('article_no')}   page_no: {g('page_no')}")
    print(f"  verify_status: {g('verify_status')}   status: {g('status')}   timeliness: {g('timeliness')}")
    print(f"  review_note  : {g('review_note')}")
    print(f"  quote_text   : {str(g('quote_text'))[:400]}")
    print(f"  content      : {str(g('content'))[:900]}")
