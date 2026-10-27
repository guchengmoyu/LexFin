# -*- coding: utf-8 -*-
"""审计第四轮：完整条文缺页、keywords/difficulty 等字段质量、同条包含关系对详情、悬空指代精确判定"""
import openpyxl, re, collections

XLSX = r'成员1_知识点审核.xlsx'
wb = openpyxl.load_workbook(XLSX, read_only=True)
ws = wb['知识点候选']
hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
I = {h: i for i, h in enumerate(hdr)}
rows = list(ws.iter_rows(min_row=2, values_only=True))
arts = {}
for r in wb['完整条文'].iter_rows(min_row=2, values_only=True):
    arts[(r[0], r[1])] = {'text': r[2] or '', 'page': r[4]}

def norm(s):
    return re.sub(r'[\s\u3000]', '', str(s or ''))

print('===== 1. 完整条文中 page 缺失的条文 =====')
for k, v in arts.items():
    if v['page'] in (None, ''):
        print(k, '-> page缺失, text:', v['text'][:50])

print()
print('===== 2. 字段填充率 =====')
for h in hdr:
    n_empty = sum(1 for r in rows if r[I[h]] in (None, ''))
    print(f'{h:20s} 空={n_empty}')

print()
print('===== 3. keywords 样例（前8行 + 疑似乱码计数） =====')
bad_kw = 0
for r in rows[:8]:
    print(r[I['code']], '|', str(r[I['keywords']])[:70])
for r in rows:
    kw = str(r[I['keywords']] or '')
    # 乱码特征：出现长度<=2的碎片 token 且 token 数>=5
    toks = [t for t in kw.split(',') if t.strip()]
    if any(len(t.strip()) <= 2 for t in toks) and len(toks) >= 5:
        bad_kw += 1
print('keywords疑似乱码行数:', bad_kw)

print()
print('===== 4. 同条包含关系对抽样详情（判断拆分 vs 重复） =====')
pairs = [('KP-INS-0299','KP-INS-0300'), ('KP-INS-0347','KP-INS-0350'), ('KP-INS-0347','KP-INS-0351'),
         ('KP-SEC-0013','KP-SEC-0014'), ('KP-SEC-0014','KP-SEC-0015'),
         ('KP-SEC-0069','KP-SEC-0070'), ('KP-SEC-0070','KP-SEC-0071'),
         ('KP-SEC-0173','KP-SEC-0174'), ('KP-SEC-0174','KP-SEC-0175'),
         ('KP-SEC-0022','KP-SEC-0023'), ('KP-SEC-0222','KP-SEC-0224')]
d = {r[I['code']]: r for r in rows}
for a, b in pairs:
    ra, rb = d.get(a), d.get(b)
    if not ra or not rb:
        continue
    print(f"--- {a}[{ra[I['article_no']]}] vs {b}[{rb[I['article_no']]}]")
    print(' A title:', ra[I['title']])
    print(' B title:', rb[I['title']])
    print(' A quote:', str(ra[I['quote_text']])[:70])
    print(' B quote:', str(rb[I['quote_text']])[:70])

print()
print('===== 5. 悬空指代精确判定：quote 起点>0 且含“前款/第一款/上述/前条”，且指代对象不在 quote 内 =====')
for r in rows:
    a = arts.get((r[I['doc_title']], r[I['article_no']]))
    if not a:
        continue
    qn, an = norm(r[I['quote_text']]), norm(a['text'])
    pos = an.find(qn)
    if pos <= 0:
        continue
    q = str(r[I['quote_text']] or '')
    if re.search(r'前款|第一款|上述|前条', q):
        # 指代对象（条文前半部分）是否在 content 的上下文段中
        c = str(r[I['content']] or '')
        head = a['text'][:max(a['text'].find(q) - 10, 0)] if a['text'].find(q) > 0 else ''
        ctx_ok = any(norm(head[-k:]) and norm(head[-k:]) in norm(c) for k in (40, 25, 15)) if head else True
        print(f"{r[I['code']]} {r[I['article_no']]} quote起始{pos}/{len(an)} content补上下文={ctx_ok} | {q[:55]}")
