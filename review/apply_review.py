# -*- coding: utf-8 -*-
"""人工审批写回：标题修复、verify_status 裁决、page_no、review_note、文档登记版本核查"""
import openpyxl, shutil, os, re

SRC = r'成员1_知识点审核.xlsx'
BAK = r'review\成员1_知识点审核_审核前备份.xlsx'
os.makedirs(r'review', exist_ok=True)
if not os.path.exists(BAK):
    shutil.copy2(SRC, BAK)
    print('已备份 ->', BAK)

wb = openpyxl.load_workbook(SRC)
ws = wb['知识点候选']
hdr = [c.value for c in ws[1]]
I = {h: i + 1 for i, h in enumerate(hdr)}   # 1-based列号

# ---------- 1. 标题修复（截断/病句/悬空指代 -> 完整问句） ----------
NEW_TITLES = {
 'KP-SEC-0064': '上市公司持有百分之五以上股份的股东、实际控制人、董监高等转让其持有的本公司股份有哪些具体规定？',
 'KP-SEC-0069': '法定禁止参与股票交易的人员在任期或者法定限期内能否持有、买卖股票或者其他具有股权性质的证券？',
 'KP-SEC-0070': '成为法定禁止参与股票交易的人员之前已持有的股票或者其他具有股权性质的证券如何依法转让？',
 'KP-SEC-0080': '公司董事会不收回短线交易收益的，负有责任的董事承担什么责任？',
 'KP-SEC-0089': '禁止内幕信息的知情人和非法获取内幕信息的人利用内幕信息从事什么活动？',
 'KP-SEC-0096': '禁止金融机构从业人员、监管和行业协会工作人员利用未公开信息从事哪些行为？',
 'KP-SEC-0098': '操纵证券市场的手段包括哪些？',
 'KP-SEC-0101': '禁止证券服务机构及其从业人员、证券业协会及监管机构工作人员在证券交易活动中作出什么行为？',
 'KP-SEC-0112': '持股达到百分之五后投资者负有哪些报告、通知、公告义务和买卖限制？',
 'KP-SEC-0117': '持股达到百分之三十继续进行收购的应如何发出收购要约？',
 'KP-SEC-0128': '协议收购达到百分之三十继续进行收购的应如何发出收购要约？',
 'KP-SEC-0141': '发生可能对股价产生较大影响的重大事件时公司如何履行临时报告义务？',
 'KP-SEC-0154': '信息披露违法致使投资者遭受损失的由谁承担赔偿责任？',
 'KP-SEC-0177': '代表人诉讼中可能存在相同诉讼请求的其他众多投资者的，人民法院可否发出公告？',
 'KP-SEC-0204': '突发性事件导致交易结果重大异常时证券交易所可采取哪些措施？',
 'KP-SEC-0219': '国务院证券监督管理机构应在多长期限内作出证券公司设立审批决定？',
 'KP-SEC-0221': '证券公司应在多长期限内申请经营证券业务许可证？',
 'KP-SEC-0232': '证券公司的董事、监事、高级管理人员应符合哪些任职条件？',
 'KP-SEC-0266': '证券公司股东虚假出资、抽逃出资被责令改正前，监管机构可否限制其股东权利？',
 'KP-SEC-0269': '证券公司被停业整顿、托管、接管或清算期间，可对其责任人员采取哪些措施？',
 'KP-SEC-0276': '不适用全国集中统一运营方式的证券，其登记、结算可否委托其他机构办理？',
 'KP-SEC-0296': '结算参与人未按时履行交收义务的，证券登记结算机构可否处理交收财产？',
 'KP-SEC-0313': '涉嫌证券违法的当事人作出承诺后，国务院证券监督管理机构可否中止或者终止调查？',
}

# ---------- 2. verify_status 裁决与 review_note（passed 留空，只写具体问题） ----------
NOTE_CROSSREF = ('引用《公司法》第一百四十六条系2018年修正文本条号；'
                 '现行《公司法》（2023年修订，2024-07-01施行）对应第一百七十八条，需按现行条号核对。')
DECISIONS = {
 'KP-SEC-0232': ('passed', ''),                    # 标题已修复；quote 为本法原文，无跨法依赖，原 doubt 理由不适用本行
 'KP-SEC-0191': ('doubted', NOTE_CROSSREF),
 'KP-SEC-0233': ('doubted', NOTE_CROSSREF),
 'KP-INS-0381': ('doubted', '完整条文表未记录本条页码，原始PDF缺失，page_no待补。'),
 'KP-SEC-0001': ('rejected', '立法目的条款，不计入核心制度事实。'),
 'KP-SEC-0032': ('rejected', '与KP-SEC-0020重复，保留该编号。'),
}

# ---------- 3. page_no 修正 ----------
PAGE_FIX = {'KP-SEC-0032': '7'}   # 第十六条条文页（原值5系从KP-SEC-0020复制的错误）

changes = []
for row in ws.iter_rows(min_row=2):
    code = row[I['code'] - 1].value
    if code in NEW_TITLES:
        old = row[I['title'] - 1].value
        row[I['title'] - 1].value = NEW_TITLES[code]
        changes.append(f'[title] {code}: {old} -> {NEW_TITLES[code]}')
    if code in DECISIONS:
        st, note = DECISIONS[code]
        old_st = row[I['verify_status'] - 1].value
        old_note = row[I['review_note'] - 1].value
        row[I['verify_status'] - 1].value = st
        row[I['review_note'] - 1].value = note if note else None
        changes.append(f'[verify_status] {code}: {old_st} -> {st} | note: {old_note} -> {note or "(空)"}')
    if code in PAGE_FIX:
        old_p = row[I['page_no'] - 1].value
        row[I['page_no'] - 1].value = PAGE_FIX[code]
        changes.append(f'[page_no] {code}: {old_p} -> {PAGE_FIX[code]}')

# ---------- 4. 文档登记：版本核查结论（不写审核人/时间） ----------
ws3 = wb['文档登记']
h3 = {c.value: i + 1 for i, c in enumerate(ws3[1])}
VER = {
 '保险销售行为管理办法': '已核对官方发布信息：国家金融监督管理总局令2023年第2号，2023-09-20公布，2024-03-01施行，未见修订或废止。',
 '中华人民共和国证券法': '已核对国家法律法规数据库：2019-12-28修订，2020-03-01施行，未见更新版本。',
}
for row in ws3.iter_rows(min_row=2):
    t = row[h3['title'] - 1].value
    if t in VER:
        old = row[h3['review_note'] - 1].value
        row[h3['review_note'] - 1].value = VER[t]
        changes.append(f'[文档登记.review_note] {t}: {old} -> {VER[t]}')

wb.save(SRC)
print(f'共 {len(changes)} 处修改：')
for c in changes:
    print(' ', c)
