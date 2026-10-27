"""按陈玉婷复核意见修订《成员1_知识点审核》。

复核意见（2026-10-21）：
  1. KP-INS-0381：备注"原始PDF缺失"已过时；解释机关类条款优先级低，暂不纳入核心知识点。
  2. KP-SEC-0232：quote 无公司法条号，但 content 的"同条上下文"仍含旧条号，
     改 passed 前需确认，不能只看本行引用。
  3. KP-SEC-0064 / 0101 / 0269 标题缩略了适用主体或情形，需补全；
     KP-SEC-0070 标题应避免暗示"任职前必须转让"。

就地更新：成员1_知识点审核_待修订.xlsx
原始未改版本：成员1_知识点审核(1).xlsx
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "成员1_知识点审核_待修订.xlsx"

NOTE = "陈玉婷复核(10-21)"

# code -> (新标题 or None, 新 verify_status or None, 新 review_note)
FIXES = {
    "KP-INS-0381": (
        None,
        "rejected",
        f"{NOTE}：条款内容为「本办法由金融监管总局负责解释」，属解释机关条款，"
        f"非可校验的核心制度事实，与 KP-SEC-0001（立法目的条款被 rejected）同口径，"
        f"暂不纳入核心知识点。原「原始PDF缺失」备注已过时；若日后改为保留，"
        f"需依原始 PDF 补填 page_no 后再转 passed。",
    ),
    "KP-SEC-0232": (
        None,
        "doubted",
        f"{NOTE}：本行 quote_text 无《公司法》条号，但 content 的「适用条件与同条上下文」"
        f"段仍含《公司法》第一百四十六条旧条号（2018 修正文本；现行《公司法》2023 修订后"
        f"对应条号已变）。原审核仅比对 quote 即改 passed 属疏漏，与 KP-SEC-0191/0233 "
        f"处理不一致，现退回 doubted，待按现行条号统一处理上下文后再定。",
    ),
    "KP-SEC-0064": (
        "上市公司持股 5% 以上股东、实际控制人、董监高及其他限售股股东转让本公司股份有哪些具体规定？",
        None,
        f"{NOTE}：标题补全适用主体「其他持有发行人首次公开发行前发行的股份或"
        f"上市公司向特定对象发行的股份的股东」（原题以「等」字省略）。",
    ),
    "KP-SEC-0101": (
        "禁止证券交易场所、证券公司、证券登记结算与证券服务机构及其从业人员、"
        "证券业协会及监管机构工作人员作出什么行为？",
        None,
        f"{NOTE}：标题补全适用主体「证券交易场所、证券公司、证券登记结算机构」"
        f"（原题仅列证券服务机构，遗漏三类主体）。",
    ),
    "KP-SEC-0269": (
        "证券公司被停业整顿、托管、接管、清算期间或出现重大风险时，"
        "可对其责任人员采取哪些措施？",
        None,
        f"{NOTE}：标题补全适用情形「或者出现重大风险时」（原题仅列停业整顿/"
        f"托管/接管/清算四种情形）。",
    ),
    "KP-SEC-0070": (
        "成为法定禁止股票交易人员时，原已持有的股票或其他具有股权性质的证券应如何处理？",
        None,
        f"{NOTE}：原标题「成为……之前已持有的股票如何依法转让」易误读为"
        f"任职前即须转让；法条为「任何人在成为前款所列人员时」才必须依法转让。"
        f"改为以「成为……时」起句，并保留「其他具有股权性质的证券」这一标的范围。",
    ),
}


def main() -> None:
    if not XLSX.exists():
        raise SystemExit(f"找不到文件: {XLSX}")
    wb = load_workbook(XLSX)
    ws = wb["知识点候选"]
    header = [str(c.value).strip() if c.value else "" for c in ws[1]]
    col = {n: i + 1 for i, n in enumerate(header)}   # 1-based

    seen = set()
    for r in range(2, ws.max_row + 1):
        code = ws.cell(r, col["code"]).value
        if code not in FIXES:
            continue
        new_title, new_verify, new_note = FIXES[code]
        seen.add(code)
        before = (ws.cell(r, col["title"]).value, ws.cell(r, col["verify_status"]).value)
        if new_title:
            ws.cell(r, col["title"]).value = new_title
        if new_verify:
            ws.cell(r, col["verify_status"]).value = new_verify
        if new_note:
            old = ws.cell(r, col["review_note"]).value
            ws.cell(r, col["review_note"]).value = (str(old).strip() + " ｜ " if old else "") + new_note
        after = (ws.cell(r, col["title"]).value, ws.cell(r, col["verify_status"]).value)
        print(f"[第{r:>4}行] {code}")
        if before[1] != after[1]:
            print(f"    verify_status: {before[1]} -> {after[1]}")
        if before[0] != after[0]:
            print(f"    title: {str(before[0])[:30]}...")
            print(f"        -> {str(after[0])[:30]}...")

    missing = set(FIXES) - seen
    if missing:
        print(f"[警告] 未找到: {sorted(missing)}")

    try:
        wb.save(XLSX)
        out = XLSX
    except PermissionError:
        out = XLSX.with_name("成员1_知识点审核_已修订.xlsx")
        wb.save(out)
        print(f"\n[提示] 原文件被占用（可能正在 Excel 中打开），已另存为新文件。")
    print(f"\n已保存: {out.name}")


if __name__ == "__main__":
    main()
