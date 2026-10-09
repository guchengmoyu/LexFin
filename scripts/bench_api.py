#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
API 时延压测脚本（T1/T2 分位数统计）
====================================

对应设计文档：
  - docs/09-测试方案设计.md §5.1（响应时间口径 T1/T2/T3/T4）、§5.2（bench 输出格式）、§5.3（并发）
  - docs/06-后端接口设计.md §5.7（POST /retrieve / GET /search）、§5.8（/qa/ask/stream SSE 事件格式）

测量口径（docs/09 §5.1）：
  T1 检索响应时间：POST /retrieve（或 GET /search）从发出请求到**完整接收响应体**的耗时
                   目标 P95 ≤ 500 ms（P50 ≤ 300 ms）
  T2 问答首 token：POST /qa/ask/stream 从发出请求到**第一个 SSE 事件**（忽略 ping 心跳，
                   meta 或首个 token 取先到者）的耗时，目标 P95 ≤ 2000 ms（P50 ≤ 1500 ms）
  T3 完整回答时间：到 done 事件的耗时，仅作参考值（P95 ≤ 8 s，非硬性验收指标）

测量流程：预热 --warmup 次（默认 5）后，对问题列表按轮询方式连续测量 --repeat 次（默认串行），
输出控制台表格（含 ✅/❌ 达标判定）与 JSON 报告（含每次样本 + latency_detail 分阶段耗时均值）。

用法示例（docs/09 §5.2）：
  python scripts/bench_api.py \\
    --base-url http://localhost:8000/api/v1 \\
    --username admin --password Admin@123 \\
    --questions data/eval/questions.txt \\
    --repeat 100 \\
    --out reports/bench_20250601.json

常用可选参数：
  --concurrency 20       并发压测（asyncio + httpx.AsyncClient，默认 1 = 串行）
  --phase both|t1|t2     只测某一阶段
  --t1-endpoint retrieve|search   T1 走 POST /retrieve 还是 GET /search
  --per-question         每个问题各测 --repeat 次（默认 --repeat 为总采样数，问题轮询）
  --warmup 5             预热次数（不计入统计）

问题列表文件：
  - .txt：一行一个问题，空行与 # 开头的行忽略；
  - .xlsx：读取表头为 question 的列（如评测集 eval_set.xlsx）。

依赖：Python 3.11+ 标准库 + httpx（openpyxl 仅在读 xlsx 问题列表时延迟导入）。
本脚本是纯 HTTP 客户端，不 import 本项目 app 包。
"""

from __future__ import annotations

import argparse
import asyncio
import json
import math
import re
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Sequence

try:
    import httpx
except ImportError:  # pragma: no cover - 环境缺依赖时给出中文提示
    print("错误：缺少依赖 httpx，请先执行 `pip install httpx`。", file=sys.stderr)
    raise SystemExit(2)


# ---------------------------------------------------------------------------
# 通用工具
# ---------------------------------------------------------------------------


class AppError(RuntimeError):
    """脚本级可读错误（统一中文提示）。"""


def percentile(values: Sequence[float], p: float) -> float:
    """最近秩（nearest-rank）分位数，p 取 0~100；空序列返回 0.0。"""
    if not values:
        return 0.0
    ordered = sorted(values)
    idx = math.ceil(p / 100 * len(ordered)) - 1
    idx = max(0, min(len(ordered) - 1, idx))
    return round(float(ordered[idx]), 1)


def now_iso() -> str:
    """本地时区 ISO8601 时间戳。"""
    return datetime.now().astimezone().isoformat(timespec="seconds")


def str2bool(v: str | bool) -> bool:
    """把命令行的 true/false/1/0 等写法转成 bool（用于 --rerank false 这类写法）。"""
    if isinstance(v, bool):
        return v
    s = str(v).strip().lower()
    if s in ("true", "1", "yes", "y", "on"):
        return True
    if s in ("false", "0", "no", "n", "off"):
        return False
    raise argparse.ArgumentTypeError("布尔参数仅接受 true/false（例如 --rerank false）")


# ---------------------------------------------------------------------------
# SSE 解析（fetch 流式语义：按空行切事件块，解析 event:/data: 行，data 为 JSON）
# ---------------------------------------------------------------------------


@dataclass
class SSEEvent:
    """单个 SSE 事件：事件名 + 原始 data 文本 + JSON 解析结果（解析失败则为 None）。"""

    event: str
    data: str
    payload: Any | None


class SSEParser:
    """增量 SSE 解析器：喂入任意大小的字节块，吐出完整事件。

    - 事件块以空行（\\n\\n 或 \\r\\n\\r\\n）分隔；
    - 支持多行 data（按 \\n 拼接）、注释行（以 ':' 开头，如心跳注释）；
    - 未声明 event: 行时事件名默认 message。
    """

    _BLOCK_SEP = re.compile(r"\r?\n\r?\n")

    def __init__(self) -> None:
        self._buf: str = ""

    def feed(self, chunk: bytes) -> list[SSEEvent]:
        """喂入字节流片段，返回本次解析出的完整事件列表。"""
        self._buf += chunk.decode("utf-8", errors="replace")
        events: list[SSEEvent] = []
        while True:
            m = self._BLOCK_SEP.search(self._buf)
            if not m:
                break
            block, self._buf = self._buf[: m.start()], self._buf[m.end():]
            ev = self._parse_block(block)
            if ev is not None:
                events.append(ev)
        return events

    def flush(self) -> list[SSEEvent]:
        """流结束时把缓冲区里未以空行结尾的残余块也解析掉。"""
        events: list[SSEEvent] = []
        if self._buf.strip():
            ev = self._parse_block(self._buf)
            if ev is not None:
                events.append(ev)
        self._buf = ""
        return events

    @staticmethod
    def _parse_block(block: str) -> SSEEvent | None:
        event_name = "message"
        data_lines: list[str] = []
        for raw in block.splitlines():
            line = raw.rstrip("\r")
            if not line or line.startswith(":"):
                continue  # 空行/注释（心跳注释）直接忽略
            if line.startswith("event:"):
                event_name = line[len("event:"):].strip()
            elif line.startswith("data:"):
                data_lines.append(line[len("data:"):].lstrip())
        if not data_lines and event_name == "message":
            return None
        data = "\n".join(data_lines)
        payload: Any | None
        try:
            payload = json.loads(data) if data else None
        except json.JSONDecodeError:
            payload = None  # data 非 JSON 时保留原文
        return SSEEvent(event=event_name, data=data, payload=payload)


# ---------------------------------------------------------------------------
# 统计
# ---------------------------------------------------------------------------


@dataclass
class PhaseStat:
    """单阶段（T1/T2/T3）统计结果。"""

    name: str
    label: str
    samples: list[dict[str, Any]] = field(default_factory=list)
    target_p95_ms: float | None = None

    @property
    def ok_values(self) -> list[float]:
        return [float(s["ms"]) for s in self.samples if s.get("ok")]

    @property
    def failed(self) -> int:
        return sum(1 for s in self.samples if not s.get("ok"))

    def summary(self) -> dict[str, Any]:
        vals = self.ok_values
        out: dict[str, Any] = {
            "name": self.name,
            "label": self.label,
            "n": len(vals),
            "failed": self.failed,
            "p50": percentile(vals, 50),
            "p90": percentile(vals, 90),
            "p95": percentile(vals, 95),
            "p99": percentile(vals, 99),
            "max": percentile(vals, 100),
            "min": percentile(vals, 0),
            "mean": round(sum(vals) / len(vals), 1) if vals else 0.0,
            "target_p95_ms": self.target_p95_ms,
        }
        out["passed"] = (
            self.target_p95_ms is not None
            and bool(vals)
            and out["p95"] <= self.target_p95_ms
            and self.failed == 0
        )
        return out


# ---------------------------------------------------------------------------
# 问题列表
# ---------------------------------------------------------------------------


def load_questions(path: Path | None, inline: Sequence[str] | None) -> list[str]:
    """读取问题列表：.txt 一行一题 / .xlsx 读 question 列；并合并 --question 内联题目。"""
    questions: list[str] = []
    if path is not None:
        if not path.exists():
            raise AppError(f"问题列表文件不存在：{path}")
        suffix = path.suffix.lower()
        if suffix in (".txt", ".text", ".list"):
            for line in path.read_text(encoding="utf-8-sig").splitlines():
                q = line.strip()
                if q and not q.startswith("#"):
                    questions.append(q)
        elif suffix in (".xlsx", ".xlsm"):
            try:
                from openpyxl import load_workbook
            except ImportError as exc:  # pragma: no cover
                raise AppError("缺少依赖 openpyxl，请先执行 `pip install openpyxl`。") from exc
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb.active
            rows = ws.iter_rows(values_only=True)
            header = [str(c).strip() if c is not None else "" for c in next(rows, ())]
            if "question" not in header:
                raise AppError(f"xlsx 问题列表缺少 question 列：{path}")
            idx = header.index("question")
            for row in rows:
                if idx < len(row) and row[idx] is not None and str(row[idx]).strip():
                    questions.append(str(row[idx]).strip())
        else:
            raise AppError(f"不支持的问题列表格式（仅 .txt / .xlsx）：{path}")
    if inline:
        questions.extend(q.strip() for q in inline if q.strip())
    if not questions:
        raise AppError("问题列表为空：请用 --questions 指定文件，或用 --question 直接传题。")
    return questions


# ---------------------------------------------------------------------------
# 单次测量（async，供串行与并发共用）
# ---------------------------------------------------------------------------


def _unwrap_error(resp: httpx.Response, raw: str) -> str:
    """把统一响应结构里的错误信息拼成中文提示（docs/06 §2/§3）。"""
    try:
        payload = json.loads(raw)
        return (
            f"HTTP {resp.status_code}, code={payload.get('code')}：{payload.get('message')}"
            f"（request_id={payload.get('request_id')}）"
        )
    except Exception:  # noqa: BLE001
        return f"HTTP {resp.status_code}：{raw[:200]}"


async def measure_t1(
    client: httpx.AsyncClient,
    args: argparse.Namespace,
    question: str,
) -> dict[str, Any]:
    """T1：POST /retrieve 或 GET /search，计时到完整接收响应体（docs/09 §5.1）。"""
    start = time.perf_counter()
    try:
        if args.t1_endpoint == "search":
            resp = await client.get(
                f"{args.base_url}/search",
                params={"q": question, "mode": args.mode, "page": 1, "size": args.top_k, "highlight": "false"},
            )
        else:
            body: dict[str, Any] = {
                "query": question,
                "mode": args.mode,
                "top_k": args.top_k,
                "include_knowledge_points": True,
                "include_chunks": True,
            }
            if args.kb_ids:
                body["kb_ids"] = args.kb_ids
            resp = await client.post(f"{args.base_url}/retrieve", json=body)
        raw = resp.text  # 显式读完响应体再停表
        ms = (time.perf_counter() - start) * 1000.0
        if resp.status_code >= 400:
            return {"ok": False, "ms": round(ms, 1), "status_code": resp.status_code,
                    "error": _unwrap_error(resp, raw)}
        payload = json.loads(raw)
        if payload.get("code") not in (0, None):
            return {"ok": False, "ms": round(ms, 1), "status_code": resp.status_code,
                    "error": _unwrap_error(resp, raw)}
        return {"ok": True, "ms": round(ms, 1), "status_code": resp.status_code, "error": None}
    except httpx.HTTPError as exc:
        ms = (time.perf_counter() - start) * 1000.0
        return {"ok": False, "ms": round(ms, 1), "status_code": None, "error": f"网络错误：{exc}"}
    except json.JSONDecodeError:
        ms = (time.perf_counter() - start) * 1000.0
        return {"ok": False, "ms": round(ms, 1), "status_code": resp.status_code,
                "error": "响应不是合法 JSON"}


async def measure_t2(
    client: httpx.AsyncClient,
    args: argparse.Namespace,
    question: str,
) -> dict[str, Any]:
    """T2：POST /qa/ask/stream 到第一个 SSE 事件（忽略 ping）；顺带记录 T3（到 done）。"""
    body: dict[str, Any] = {
        "question": question,
        "session_no": None,
        "kb_ids": args.kb_ids,
        "categories": None,
        "timeliness": None,
        "mode": args.mode,
        "top_k": args.top_k,
        "stream": True,
        "use_history": False,
    }
    if args.rerank is not None:
        body["rerank"] = args.rerank

    start = time.perf_counter()
    parser = SSEParser()
    ttft: float | None = None
    done_ms: float | None = None
    done_payload: dict[str, Any] = {}
    meta_payload: dict[str, Any] = {}
    answer_parts: list[str] = []
    error: str | None = None
    status_code: int | None = None

    def _handle(ev: SSEEvent) -> bool:
        """处理单个事件；返回是否结束流。"""
        nonlocal ttft, done_ms, error
        if ev.event == "ping":
            return False  # 心跳忽略，不参与 TTFT
        if ttft is None:
            ttft = (time.perf_counter() - start) * 1000.0
        payload = ev.payload if isinstance(ev.payload, dict) else {}
        if ev.event == "meta":
            meta_payload.update(payload)
        elif ev.event == "token":
            answer_parts.append(str(payload.get("delta", "")))
        elif ev.event == "done":
            done_payload.update(payload)
            done_ms = (time.perf_counter() - start) * 1000.0
            return True
        elif ev.event == "error":
            error = f"服务端 error 事件（code={payload.get('code')}）：{payload.get('message')}"
            return True
        return False

    try:
        async with client.stream("POST", f"{args.base_url}/qa/ask/stream", json=body) as resp:
            status_code = resp.status_code
            if resp.status_code >= 400:
                raw = (await resp.aread()).decode("utf-8", errors="replace")
                ms = (time.perf_counter() - start) * 1000.0
                return {"ok": False, "ms": round(ms, 1), "status_code": resp.status_code,
                        "error": _unwrap_error(resp, raw)}
            finished = False
            async for chunk in resp.aiter_bytes():
                for ev in parser.feed(chunk):
                    if _handle(ev):
                        finished = True
                        break
                if finished:
                    break
            for ev in parser.flush():
                _handle(ev)
    except httpx.HTTPError as exc:
        ms = (time.perf_counter() - start) * 1000.0
        return {"ok": False, "ms": round(ms, 1), "status_code": status_code, "error": f"网络错误：{exc}"}

    if ttft is None:
        return {"ok": False, "ms": round((time.perf_counter() - start) * 1000.0, 1),
                "status_code": status_code,
                "error": error or "未收到任何 SSE 事件（meta/token/done/error）"}
    if error:
        return {"ok": False, "ms": round(ttft, 1), "status_code": status_code, "error": error}

    answer = "".join(answer_parts)
    used = done_payload.get("used_citations")
    has_citation = bool(re.search(r"\[\d+\]", answer)) or bool(used) or bool(meta_payload.get("citations"))
    return {
        "ok": True,
        "ms": round(ttft, 1),  # T2 = 首个 SSE 事件时刻
        "status_code": status_code,
        "error": None,
        "t3_ms": round(done_ms, 1) if done_ms is not None else None,
        "answer_type": done_payload.get("answer_type"),
        "has_citation": has_citation,
        "latency_detail": done_payload.get("latency_detail") or {},
        "latency_ms": done_payload.get("latency_ms"),
    }


# ---------------------------------------------------------------------------
# 压测主流程
# ---------------------------------------------------------------------------


async def run_bench(args: argparse.Namespace, questions: list[str]) -> dict[str, Any]:
    """预热 + 连续测量（串行或并发），返回 JSON 报告 dict。"""
    t1 = PhaseStat("t1", "T1 检索响应时间", target_p95_ms=args.t1_target_ms)
    t2 = PhaseStat("t2", "T2 问答首 token", target_p95_ms=args.t2_target_ms)
    t3 = PhaseStat("t3", "T3 问答完整响应（参考）", target_p95_ms=None)
    total = args.repeat * (len(questions) if args.per_question else 1)

    timeout = httpx.Timeout(args.timeout, read=args.timeout)
    async with httpx.AsyncClient(
        timeout=timeout,
        headers={"User-Agent": "FinRAG-bench_api/1.0"},
        follow_redirects=False,
    ) as client:
        # ---- 登录拿 token（docs/06 §5.1） ----
        token = args.token
        if not token:
            if not args.username or not args.password:
                raise AppError("缺少凭证：请提供 --username/--password 或 --token。")
            resp = await client.post(
                f"{args.base_url}/auth/login",
                json={"username": args.username, "password": args.password},
            )
            try:
                payload = resp.json()
            except Exception as exc:  # noqa: BLE001
                raise AppError(f"登录响应不是合法 JSON（HTTP {resp.status_code}）") from exc
            if resp.status_code >= 400 or payload.get("code") not in (0, None):
                raise AppError(f"登录失败：{_unwrap_error(resp, resp.text)}")
            token = (payload.get("data") or {}).get("access_token")
            if not token:
                raise AppError("登录响应中未找到 data.access_token。")
        client.headers["Authorization"] = f"Bearer {token}"

        def want_t1() -> bool:
            return args.phase in ("both", "t1")

        def want_t2() -> bool:
            return args.phase in ("both", "t2")

        # ---- 预热（默认 5 次，不计入统计，docs/09 §5.1 测量条件） ----
        if args.warmup > 0:
            print(f"预热中（{args.warmup} 次，不计入统计）……")
            for i in range(args.warmup):
                q = questions[i % len(questions)]
                if want_t1():
                    await measure_t1(client, args, q)
                if want_t2():
                    await measure_t2(client, args, q)
                if args.interval > 0:
                    await asyncio.sleep(args.interval)

        # ---- 测量：并发模式用信号量控制在飞请求数 ----
        print(
            f"开始测量：总样本 {total} 次/阶段   问题数 {len(questions)}   "
            f"并发数 {args.concurrency}（1 = 串行）   阶段 {args.phase}"
        )
        sem = asyncio.Semaphore(max(1, args.concurrency))
        progress = {"done": 0}
        lock = asyncio.Lock()

        async def one(seq: int, question: str) -> None:
            async with sem:
                if want_t1():
                    res = await measure_t1(client, args, question)
                    res.update({"seq": seq, "question": question})
                    t1.samples.append(res)
                if want_t2():
                    res2 = await measure_t2(client, args, question)
                    res2.update({"seq": seq, "question": question})
                    t2.samples.append(res2)
                    if res2.get("t3_ms") is not None:
                        t3.samples.append({"seq": seq, "question": question, "ok": True,
                                           "ms": res2["t3_ms"], "error": None})
                async with lock:
                    progress["done"] += 1
                    if progress["done"] % 10 == 0 or progress["done"] == total:
                        print(f"  进度 {progress['done']}/{total}")
                if args.interval > 0:
                    await asyncio.sleep(args.interval)

        if args.per_question:
            # 每题各测 --repeat 次：序号 i 对应第 i // repeat 题
            jobs = [(i, questions[(i // args.repeat) % len(questions)]) for i in range(total)]
        else:
            # 默认：--repeat 为总采样数，问题按列表轮询（docs/09 §5.2 示例 n=100）
            jobs = [(i, questions[i % len(questions)]) for i in range(total)]
        await asyncio.gather(*(one(seq, q) for seq, q in jobs))

    # ---- 分阶段耗时均值：取 done 事件 latency_detail（docs/06 §5.8） ----
    stage_keys = ("rewrite", "retrieve", "rerank", "llm_ttft", "prompt", "network", "total")
    stage_sums: dict[str, list[float]] = {}
    for s in t2.samples:
        detail = s.get("latency_detail") or {}
        for k, v in detail.items():
            if isinstance(v, (int, float)):
                stage_sums.setdefault(k, []).append(float(v))
    stages_mean = {
        k: round(sum(v) / len(v), 1)
        for k, v in sorted(stage_sums.items())
        if v
    }
    ttfts = t2.ok_values
    stages_mean["total_ttft"] = round(sum(ttfts) / len(ttfts), 1) if ttfts else 0.0
    # 缺失的阶段键补 0，保证报告字段稳定（rewrite/retrieve/rerank/llm_ttft/total）
    for k in stage_keys:
        stages_mean.setdefault(k, 0.0)

    # ---- 回答质量抽样（来自 T2 的 done 事件） ----
    answered = sum(1 for s in t2.samples if s.get("answer_type") == "answered")
    refused = sum(1 for s in t2.samples if s.get("answer_type") == "refused")
    cited = sum(1 for s in t2.samples if s.get("has_citation") and s.get("answer_type") == "answered")
    quality = {
        "samples": len(t2.samples),
        "answered": answered,
        "refused": refused,
        "citation_rate": round(cited / answered, 3) if answered else None,
    }

    errors = [
        {"phase": ph.name, "seq": s.get("seq"), "question": s.get("question"),
         "status_code": s.get("status_code"), "error": s.get("error")}
        for ph in (t1, t2)
        for s in ph.samples
        if not s.get("ok")
    ]

    report = {
        "generated_at": now_iso(),
        "config": {
            "base_url": args.base_url,
            "questions_file": str(args.questions) if args.questions else None,
            "questions": len(questions),
            "repeat": args.repeat,
            "per_question": args.per_question,
            "warmup": args.warmup,
            "concurrency": args.concurrency,
            "phase": args.phase,
            "t1_endpoint": args.t1_endpoint,
            "mode": args.mode,
            "top_k": args.top_k,
            "rerank": args.rerank,
            "kb_ids": args.kb_ids,
            "timeout_s": args.timeout,
        },
        "t1": t1.summary(),
        "t2": t2.summary(),
        "t3_reference": t3.summary(),
        "stages_mean": stages_mean,
        "quality": quality,
        "errors": errors[:50],
        "samples": {"t1": t1.samples, "t2": t2.samples, "t3": t3.samples},
    }
    _print_report(report)
    return report


def _print_report(report: dict[str, Any]) -> None:
    """控制台打印 docs/09 §5.2 示例格式的压测报告。"""

    def block(title: str, st: dict[str, Any], target: float | None) -> None:
        print(f"================ {title} (n={st['n']}) ================")
        if st["n"] == 0:
            print("  （无成功样本，无法统计）")
            print()
            return
        print(
            f"  P50={st['p50']:5.0f} ms   P90={st['p90']:5.0f} ms   P95={st['p95']:5.0f} ms"
            f"   P99={st['p99']:5.0f} ms   max={st['max']:5.0f} ms"
        )
        if target is not None:
            ok = st["p95"] <= target and st["failed"] == 0
            print(f"  目标 P95 <= {target:.0f} ms  →  {'✅ 达标' if ok else '❌ 未达标'}")
        else:
            print(f"  （参考值，非硬性验收指标；P95 ≤ 8000 ms {'✅' if st['p95'] <= 8000 else '❌'}）")
        if st["failed"]:
            print(f"  ⚠️ 失败样本 {st['failed']} 个（不计入分位数，详见 JSON errors）")
        print()

    t1, t2, t3 = report["t1"], report["t2"], report["t3_reference"]
    if report["config"]["phase"] in ("both", "t1"):
        block("T1 检索响应时间", t1, t1["target_p95_ms"])
    if report["config"]["phase"] in ("both", "t2"):
        block("T2 问答首 token", t2, t2["target_p95_ms"])

    print("================ 分阶段耗时均值 ================")
    st = report["stages_mean"]
    print(
        f"  rewrite={st.get('rewrite', 0):5.0f} ms   retrieve={st.get('retrieve', 0):5.0f} ms"
        f"   rerank={st.get('rerank', 0):5.0f} ms   llm_ttft={st.get('llm_ttft', 0):5.0f} ms"
    )
    print(
        f"  prompt={st.get('prompt', 0):5.0f} ms   network={st.get('network', 0):5.0f} ms"
        f"   total_ttft={st.get('total_ttft', 0):5.0f} ms"
    )
    print("  （取 /qa/ask/stream done 事件 latency_detail 字段；network 未提供时为 0）")
    print()

    q = report["quality"]
    print("================ 回答质量抽样 ================")
    rate = q["citation_rate"]
    rate_text = f"{rate * 100:.0f}%" if rate is not None else "—"
    print(
        f"  样本数 {q['samples']}，answered {q['answered']}，refused {q['refused']}，"
        f"带引用率 {rate_text}"
    )
    print()

    print("================ 达标结论 ================")
    lines = []
    if report["config"]["phase"] in ("both", "t1"):
        lines.append(
            f"  T1 检索 P95 = {t1['p95']:.0f} ms（目标 ≤ {t1['target_p95_ms']:.0f} ms）"
            f"  →  {'✅' if t1['passed'] else '❌'}"
        )
    if report["config"]["phase"] in ("both", "t2"):
        lines.append(
            f"  T2 问答首 token P95 = {t2['p95']:.0f} ms（目标 ≤ {t2['target_p95_ms']:.0f} ms）"
            f"  →  {'✅' if t2['passed'] else '❌'}"
        )
    if report["config"]["phase"] in ("both", "t2") and t3["n"]:
        lines.append(f"  T3 完整回答 P95 = {t3['p95']:.0f} ms（参考 ≤ 8000 ms，非硬性指标）")
    lines.append(f"  并发数 = {report['config']['concurrency']}，失败样本 = {len(report['errors'])}")
    print("\n".join(lines))
    print()


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        prog="bench_api.py",
        description="API 时延压测脚本（docs/09 §5.2）：T1 检索响应时间 / T2 问答首 token 的分位数统计与达标判定",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "示例：\n"
            "  python scripts/bench_api.py --base-url http://localhost:8000/api/v1 \\\n"
            "    --username admin --password Admin@123 \\\n"
            "    --questions data/eval/questions.txt --repeat 100 --out reports/bench_20250601.json"
        ),
    )
    parser.add_argument("--base-url", required=True, help="接口基地址，如 http://localhost:8000/api/v1")
    parser.add_argument("--username", default=None, help="登录用户名（与 --token 二选一）")
    parser.add_argument("--password", default=None, help="登录密码")
    parser.add_argument("--token", default=None, help="直接提供 access_token，跳过登录")
    parser.add_argument("--questions", type=Path, default=None,
                        help="问题列表文件（.txt 一行一题 / .xlsx 读 question 列）")
    parser.add_argument("--question", action="append", default=None,
                        help="直接追加单个问题（可重复），与 --questions 可同时使用")
    parser.add_argument("--repeat", type=int, default=100,
                        help="每阶段测量次数（默认总采样数，问题按列表轮询）；--per-question 时为每题次数")
    parser.add_argument("--out", type=Path, default=None,
                        help="JSON 报告输出路径，默认 reports/bench_<时间戳>.json")
    parser.add_argument("--warmup", type=int, default=5, help="预热次数，不计入统计，默认 5")
    parser.add_argument("--concurrency", type=int, default=1,
                        help="并发数（asyncio + httpx.AsyncClient），默认 1 = 串行")
    parser.add_argument("--phase", choices=("both", "t1", "t2"), default="both",
                        help="测量阶段：both / 只测 T1 / 只测 T2，默认 both")
    parser.add_argument("--t1-endpoint", choices=("retrieve", "search"), default="retrieve",
                        help="T1 走 POST /retrieve 还是 GET /search，默认 retrieve")
    parser.add_argument("--per-question", action="store_true",
                        help="每个问题各测 --repeat 次（默认 --repeat 为总采样数，问题轮询）")
    parser.add_argument("--mode", choices=("keyword", "vector", "hybrid"), default="hybrid",
                        help="检索/问答模式，默认 hybrid")
    parser.add_argument("--top-k", type=int, default=5, help="检索 Top-K（1~10），默认 5")
    parser.add_argument("--rerank", nargs="?", const=True, default=None, type=str2bool,
                        help="重排开关，如 --rerank false；不传则由服务端默认配置决定")
    parser.add_argument("--kb-ids", default=None, help="限定知识库 id，逗号分隔，如 1,2")
    parser.add_argument("--interval", type=float, default=0.0,
                        help="样本之间的间隔秒数（串行时可缓解服务端压力），默认 0")
    parser.add_argument("--timeout", type=float, default=60.0, help="单请求超时秒数，默认 60")
    parser.add_argument("--t1-target-ms", type=float, default=500.0,
                        help="T1 P95 目标（ms），默认 500（docs/09 §5.1）")
    parser.add_argument("--t2-target-ms", type=float, default=2000.0,
                        help="T2 P95 目标（ms），默认 2000（docs/09 §5.1）")
    args = parser.parse_args(argv)

    if not args.token and (not args.username or not args.password):
        parser.error("必须提供 --username/--password 或 --token 之一")
    if args.repeat < 1:
        parser.error("--repeat 必须 ≥ 1")
    if args.warmup < 0:
        parser.error("--warmup 不能为负数")
    if args.concurrency < 1:
        parser.error("--concurrency 必须 ≥ 1")
    if not (1 <= args.top_k <= 10):
        parser.error("--top-k 取值范围为 1~10（docs/06 §5.8 参数校验）")
    if args.kb_ids:
        try:
            args.kb_ids = [int(x) for x in args.kb_ids.split(",") if x.strip()]
        except ValueError:
            parser.error("--kb-ids 需为逗号分隔的整数，如 1,2")
    else:
        args.kb_ids = None
    args.base_url = args.base_url.rstrip("/")
    return args


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        questions = load_questions(args.questions, args.question)
        report = asyncio.run(run_bench(args, questions))
        out_path = args.out or Path("reports") / f"bench_{datetime.now():%Y%m%d_%H%M%S}.json"
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"JSON 报告已写入：{out_path}")
        ok = all(
            st["passed"]
            for st, want in ((report["t1"], "t1"), (report["t2"], "t2"))
            if args.phase in ("both", want)
        )
        return 0 if ok and not report["errors"] else 1
    except KeyboardInterrupt:
        print("\n已中断（KeyboardInterrupt）。", file=sys.stderr)
        return 130
    except AppError as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
